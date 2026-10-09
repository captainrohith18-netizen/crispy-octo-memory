from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document


def build_demo_text() -> str:
    return """College Workshop Syllabus

Course: Applied RAG and Local LLMs
Instructor: Dr. Aisha Raman

Overview
This workshop introduces students to local document intelligence using Retrieval-Augmented Generation (RAG). Students will learn how to extract text from PDFs and DOCX files, split documents into meaningful chunks, generate embeddings locally, and compare query embeddings to document chunks using cosine similarity.

Learning objectives
1. Understand how document text is extracted and cleaned.
2. Learn why chunking matters for retrieval and model context.
3. Generate local embeddings with a lightweight sentence-transformer model.
4. Use cosine similarity to find the most relevant chunks.
5. Retrieve context and send it to a local Ollama model.
6. Explain why grounded answers are more reliable than unsupported guesses.

Eligibility requirements
Students should have basic Python knowledge and a willingness to experiment with local AI tools. Prior experience with notebooks is helpful but not required. Students should bring a laptop with sufficient storage for a lightweight local model and the required Python packages.

Assessment policy
The workshop includes a practical lab, a short reflection, and a final demonstration. Students are expected to complete a short Q&A exercise using a sample syllabus or document. The final activity checks whether the student can explain the retrieval flow clearly and identify why the LLM must answer using the retrieved context.

Deadlines
- Registration closes on Friday.
- Lab setup must be completed before the first practical session.
- Demo questions should be submitted at least one day before the final showcase.

Important reminder
The model is not allowed to invent facts. Answers must be grounded in the retrieved document context. If the information is missing from the uploaded document, the system should say so clearly instead of guessing.
"""


def create_demo_docx(output_path: Path) -> None:
    doc = Document()
    doc.add_heading("College Workshop Syllabus", level=1)
    for paragraph in build_demo_text().split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def create_demo_pdf(output_path: Path) -> None:
    text = build_demo_text()
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text, fontsize=11, fontname="helv")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    doc.close()


def main() -> None:
    demo_dir = Path(__file__).resolve().parents[1] / "data" / "demo"
    docx_path = demo_dir / "college_syllabus_demo.docx"
    pdf_path = demo_dir / "college_syllabus_demo.pdf"
    create_demo_docx(docx_path)
    create_demo_pdf(pdf_path)
    print(f"Created DOCX demo: {docx_path}")
    print(f"Created PDF demo: {pdf_path}")


if __name__ == "__main__":
    main()
